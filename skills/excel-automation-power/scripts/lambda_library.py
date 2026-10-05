#!/usr/bin/env python3
"""يكتب مصنّف إكسل فيه مكتبة دوال LAMBDA مسمّاة (أسماء معرّفة) + ورقة Docs تشرح كل دالة ومعاملاتها ومثالاً + ورقة Try للتجربة.

usage: python lambda_library.py --out lambda-library.xlsx [--only SAFE_DIV,PCT_CHANGE]
الدوال: SAFE_DIV, PCT_CHANGE, CAGR, TRIM_ALL, ARABIC_DIGITS, LATIN_DIGITS, BUCKET, CLAMP, NPV_DATED, WEEKDAY_AR, HIJRI_TEXT, SPLIT_NAME, LAST_NONBLANK, RUNNING_TOTAL, RANK_DENSE, IS_BETWEEN, ROUND_TO, MONTH_AR
تعمل في Excel 365/2024 (LAMBDA + LET + دوال المصفوفات). يحتاج openpyxl.
"""
import argparse
import json
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.workbook.defined_name import DefinedName
except ImportError:
    sys.exit("openpyxl غير مثبت: pip install openpyxl")

LIB = {
    "SAFE_DIV": ("=LAMBDA(num,den,[alt],LET(a,IF(ISOMITTED(alt),\"-\",alt),IF(OR(den=0,den=\"\"),a,num/den)))", "قسمة آمنة: تعيد «-» (أو بديلاً) عند قاسم صفر/فارغ", "num, den, [alt]", "=SAFE_DIV(A2,B2)"),
    "PCT_CHANGE": ("=LAMBDA(old,new,IF(OR(old=0,old=\"\"),\"\",(new-old)/ABS(old)))", "نسبة التغيّر من old إلى new مع حماية الصفر", "old, new", "=PCT_CHANGE(B2,C2)"),
    "CAGR": ("=LAMBDA(start,end,years,IF(OR(start<=0,years<=0),\"\",(end/start)^(1/years)-1))", "معدل النمو السنوي المركّب", "start, end, years", "=CAGR(B2,F2,4)"),
    "TRIM_ALL": ("=LAMBDA(txt,TRIM(SUBSTITUTE(SUBSTITUTE(CLEAN(txt),CHAR(160),\" \"),\"ـ\",\"\")))", "تنظيف النص: مسافات زائدة وغير قابلة للكسر والكشيدة", "txt", "=TRIM_ALL(A2)"),
    "ARABIC_DIGITS": ("=LAMBDA(txt,LET(t,txt&\"\",REDUCE(t,SEQUENCE(10,1,0),LAMBDA(acc,i,SUBSTITUTE(acc,TEXT(i,\"0\"),UNICHAR(1632+i))))))", "يحوّل الأرقام اللاتينية 0-9 إلى عربية-هندية ٠-٩ (نص)", "txt", "=ARABIC_DIGITS(A2)"),
    "LATIN_DIGITS": ("=LAMBDA(txt,LET(t,txt&\"\",REDUCE(t,SEQUENCE(10,1,0),LAMBDA(acc,i,SUBSTITUTE(acc,UNICHAR(1632+i),TEXT(i,\"0\"))))))", "يحوّل ٠-٩ إلى 0-9 (ثم VALUE للتحويل إلى رقم)", "txt", "=VALUE(LATIN_DIGITS(A2))"),
    "BUCKET": ("=LAMBDA(x,edges,labels,INDEX(labels,MATCH(x,edges,1)))", "تصنيف قيمة إلى شريحة: edges حدود تصاعدية وlabels أسماء الشرائح", "x, edges, labels", "=BUCKET(A2,{0,1000,5000},{\"صغير\",\"متوسط\",\"كبير\"})"),
    "CLAMP": ("=LAMBDA(x,lo,hi,MIN(MAX(x,lo),hi))", "حصر القيمة بين حدّين", "x, lo, hi", "=CLAMP(A2,0,100)"),
    "NPV_DATED": ("=LAMBDA(rate,cashflows,dates,LET(t0,MIN(dates),SUMPRODUCT(cashflows/(1+rate)^((dates-t0)/365))))", "صافي القيمة الحالية بتواريخ فعلية (XNPV بلا الحاجة لترتيب)", "rate, cashflows, dates", "=NPV_DATED(0.1,C2:C10,B2:B10)"),
    "WEEKDAY_AR": ("=LAMBDA(d,CHOOSE(WEEKDAY(d,1),\"الأحد\",\"الاثنين\",\"الثلاثاء\",\"الأربعاء\",\"الخميس\",\"الجمعة\",\"السبت\"))", "اسم اليوم بالعربية", "d", "=WEEKDAY_AR(A2)"),
    "MONTH_AR": ("=LAMBDA(d,CHOOSE(MONTH(d),\"يناير\",\"فبراير\",\"مارس\",\"أبريل\",\"مايو\",\"يونيو\",\"يوليو\",\"أغسطس\",\"سبتمبر\",\"أكتوبر\",\"نوفمبر\",\"ديسمبر\"))", "اسم الشهر بالعربية", "d", "=MONTH_AR(A2)"),
    "HIJRI_TEXT": ("=LAMBDA(d,[fmt],TEXT(d,\"[$-ar-SA,16]\"&IF(ISOMITTED(fmt),\"dd/mm/yyyy\",fmt)))", "التاريخ الهجري (أم القرى) كنص؛ fmt اختياري", "d, [fmt]", "=HIJRI_TEXT(A2)"),
    "SPLIT_NAME": ("=LAMBDA(full,part,LET(p,TEXTSPLIT(TRIM(full),\" \"),n,COUNTA(p),IF(part=\"first\",INDEX(p,1),IF(part=\"last\",INDEX(p,n),TEXTJOIN(\" \",TRUE,INDEX(p,SEQUENCE(1,MAX(1,n-2),2)))))))", "جزء من الاسم: first/last/middle", "full, part", "=SPLIT_NAME(A2,\"first\")"),
    "LAST_NONBLANK": ("=LAMBDA(rng,LET(v,FILTER(rng,rng<>\"\",\"\"),INDEX(v,ROWS(v))))", "آخر قيمة غير فارغة في عمود", "rng", "=LAST_NONBLANK(A:A)"),
    "RUNNING_TOTAL": ("=LAMBDA(rng,SCAN(0,rng,LAMBDA(acc,x,acc+N(x))))", "مجموع تراكمي (مصفوفة ديناميكية)", "rng", "=RUNNING_TOTAL(B2:B20)"),
    "RANK_DENSE": ("=LAMBDA(x,rng,SUMPRODUCT((UNIQUE(rng)>x)*1)+1)", "ترتيب كثيف (بلا فجوات) تنازلياً", "x, rng", "=RANK_DENSE(B2,$B$2:$B$50)"),
    "IS_BETWEEN": ("=LAMBDA(x,lo,hi,AND(x>=lo,x<=hi))", "هل القيمة ضمن المدى (شامل)", "x, lo, hi", "=IS_BETWEEN(A2,1,10)"),
    "ROUND_TO": ("=LAMBDA(x,step,MROUND(x,step))", "تقريب إلى أقرب خطوة (5، 0.25…)", "x, step", "=ROUND_TO(A2,0.25)"),
}


def main() -> None:
    ap = argparse.ArgumentParser(); ap.add_argument("--out", default="lambda-library.xlsx"); ap.add_argument("--only", default="")
    a = ap.parse_args()
    keys = [k for k in LIB if not a.only or k in a.only.split(",")]
    wb = Workbook(); ws = wb.active; ws.title = "Docs"; ws.sheet_view.rightToLeft = True
    ws.append(["الدالة", "الوصف", "المعاملات", "مثال", "التعريف"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3864")
    for k in keys:
        f, desc, params, ex, = LIB[k]
        ws.append([k, desc, params, ex, f])
        wb.defined_names.add(DefinedName(k, attr_text=f))
    for col, w in zip("ABCDE", (18, 60, 22, 40, 90)):
        ws.column_dimensions[col].width = w
    t = wb.create_sheet("Try"); t.sheet_view.rightToLeft = True
    t.append(["المدخل", "النتيجة", "ملاحظة"])
    t.append([1234.5, "=ARABIC_DIGITS(A2)", "أرقام عربية"]); t.append([0, "=SAFE_DIV(10,A3)", "قسمة آمنة"]); t.append([100, "=PCT_CHANGE(A4,150)", "+50%"])
    t.append(["  أحمد   محمد  ", "=TRIM_ALL(A5)", "تنظيف"]); t.append([45000, "=BUCKET(A6,{0,1000,5000},{\"صغير\",\"متوسط\",\"كبير\"})", "شريحة"])
    wb.save(a.out)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    print(json.dumps({"out": a.out, "functions": keys, "note": "افتح في Excel 365/2024؛ الأسماء في Formulas → Name Manager؛ انسخها إلى أي مصنّف عبر Name Manager أو الصق تعريفها"}, ensure_ascii=False))


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
"""يفحص ملف .xlsx: الصيغ، وأخطاء القيم المحفوظة (#REF!، #N/A…)، والأرقام المخزّنة كنص، والقيم الثابتة وسط أعمدة الصيغ، والأوراق والجداول.

usage: python xlsx_check.py workbook.xlsx [--json]
يحتاج openpyxl (pip install openpyxl). لا يعيد الحساب (إكسل يفعل ذلك عند الفتح)؛ يفحص ما هو محفوظ في الملف.
"""
import argparse
import json
import re
import sys

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("openpyxl غير مثبت: pip install openpyxl")

ERR = {"#REF!", "#N/A", "#VALUE!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!"}
VOLATILE = r"\b(NOW|TODAY|RAND|RANDBETWEEN|OFFSET|INDIRECT)\s*\("
LEGACY = {"VLOOKUP": "XLOOKUP", "HLOOKUP": "XLOOKUP", "SUMPRODUCT((": "SUMIFS/FILTER", "IFERROR(VLOOKUP": "XLOOKUP(...,if_not_found)"}


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("file")
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args()
    wb_f = load_workbook(a.file, data_only=False)
    wb_v = load_workbook(a.file, data_only=True)
    report = {"file": a.file, "sheets": [], "issues": []}
    for ws in wb_f.worksheets:
        wv = wb_v[ws.title]
        formulas, errors, text_numbers, hardcoded, volatile, legacy = 0, [], [], [], [], {}
        col_formula_rows = {}
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None:
                    continue
                if isinstance(v, str) and v.startswith("="):
                    formulas += 1
                    col_formula_rows.setdefault(c.column, []).append(c.row)
                    if re.search(VOLATILE, v, re.I):
                        volatile.append(c.coordinate)
                    for k, rep in LEGACY.items():
                        if k in v.upper():
                            legacy[k] = rep
                    cached = wv[c.coordinate].value
                    if isinstance(cached, str) and cached in ERR:
                        errors.append({"cell": c.coordinate, "error": cached, "formula": v[:80]})
                elif isinstance(v, str) and re.fullmatch(r"\s*-?[\d,]+(\.\d+)?\s*", v):
                    text_numbers.append(c.coordinate)
        # hard-coded values inside formula columns (a number sitting among formulas)
        for col, rows in col_formula_rows.items():
            if len(rows) < 4:
                continue
            lo, hi = min(rows), max(rows)
            for r in range(lo, hi + 1):
                cell = ws.cell(row=r, column=col)
                if isinstance(cell.value, (int, float)):
                    hardcoded.append(cell.coordinate)
        tables = list(getattr(ws, "tables", {}).keys())
        info = {"sheet": ws.title, "dims": ws.dimensions, "formulas": formulas, "tables": tables,
                "errors": errors[:20], "text_numbers": text_numbers[:20], "hardcoded_in_formula_columns": hardcoded[:20],
                "volatile": volatile[:10], "legacy_functions": legacy}
        report["sheets"].append(info)
        if errors:
            report["issues"].append(f"{ws.title}: {len(errors)} خلية بخطأ محفوظ (مثل {errors[0]['cell']} {errors[0]['error']})")
        if text_numbers:
            report["issues"].append(f"{ws.title}: {len(text_numbers)} رقم مخزّن كنص (لن تُجمَع) مثل {text_numbers[0]}")
        if hardcoded:
            report["issues"].append(f"{ws.title}: {len(hardcoded)} قيمة ثابتة وسط عمود صيغ مثل {hardcoded[0]}")
        if legacy:
            report["issues"].append(f"{ws.title}: دوال قديمة {list(legacy)} ← الأحدث: {list(legacy.values())}")
        if volatile:
            report["issues"].append(f"{ws.title}: دوال متقلبة ({len(volatile)}) تبطئ الملف وتغيّر النتائج")
    report["verdict"] = "PASS" if not any("خطأ محفوظ" in i or "كنص" in i for i in report["issues"]) else "REVISE"
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if a.json:
        print(json.dumps(report, ensure_ascii=False, indent=1))
    else:
        print(f"{report['verdict']}  {a.file}")
        for s in report["sheets"]:
            print(f"- {s['sheet']} {s['dims']}: {s['formulas']} صيغة، جداول: {s['tables'] or 'لا'}")
        for i in report["issues"]:
            print("  !", i)
    sys.exit(0 if report["verdict"] == "PASS" else 1)


if __name__ == "__main__":
    main()

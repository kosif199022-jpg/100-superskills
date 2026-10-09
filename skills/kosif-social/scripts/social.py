#!/usr/bin/env python3
"""KOSIF Social — the social-media side of a video or a campaign, measured, never published.

    python social.py limits [--platform reels]                 # each platform's caption/title/hashtag/length rules
    python social.py pack VIDEO --platforms tiktok,reels,shorts --out pack.json   # a post-pack skeleton from a real file
    python social.py check pack.json                           # lint a filled pack: lengths, hashtags, hook, video fit
    python social.py carousel spec.json --out DIR              # Arabic-first carousel slides (PNG + PDF + contact sheet)
    python social.py calendar plan.json --out DIR              # content calendar (CSV, XLSX, Markdown, ICS)
    python social.py besttime export.csv [--metric views]      # best weekday/hour from the account's OWN export
    python social.py abplan spec.json --out DIR                # an A/B test plan with the sample size it needs
    python social.py abread results.csv                        # did variant B really beat A? (two-proportion z-test)

Everything is local and deterministic. Nothing here posts, schedules on a platform or signs in anywhere.
Platform rules are the commonly published ones (see AS_OF); verify before a launch.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import math
import re
import shutil
import statistics
import subprocess
import sys
import unicodedata
from pathlib import Path

for _s in (sys.stdout, sys.stderr):
    if hasattr(_s, "reconfigure"):
        _s.reconfigure(encoding="utf-8", errors="replace")

AS_OF = "2026-10 — من القواعد المعلنة للمنصات؛ راجعها قبل حملة مهمة"

# caption = the post text limit (description on YouTube); visible = characters shown before "more";
# hashtags = hard maximum (None = no hard limit); seconds = longest upload treated as this format.
PLATFORMS: dict[str, dict] = {
    "tiktok":    {"ar": "تيك توك", "caption": 4000, "title": None, "hashtags": None, "hashtags_tip": "3–5",
                  "visible": None, "seconds": 600, "ratio": "9:16", "size": "1080x1920", "carousel": 35,
                  "links_clickable": False},
    "reels":     {"ar": "إنستغرام ريلز", "caption": 2200, "title": None, "hashtags": 5, "hashtags_tip": "3–5",
                  "visible": 125, "seconds": 180, "ratio": "9:16", "size": "1080x1920", "carousel": None,
                  "links_clickable": False},
    "instagram": {"ar": "إنستغرام (بوست/كاروسيل)", "caption": 2200, "title": None, "hashtags": 5, "hashtags_tip": "3–5",
                  "visible": 125, "seconds": None, "ratio": "4:5", "size": "1080x1350", "carousel": 20,
                  "links_clickable": False},
    "shorts":    {"ar": "يوتيوب شورتس", "caption": 5000, "title": 100, "hashtags": 60, "hashtags_tip": "3 (أول 3 تظهر)",
                  "visible": None, "seconds": 180, "ratio": "9:16", "size": "1080x1920", "carousel": None,
                  "links_clickable": True},
    "youtube":   {"ar": "يوتيوب", "caption": 5000, "title": 100, "hashtags": 60, "hashtags_tip": "3 (أول 3 تظهر)",
                  "visible": 150, "seconds": None, "ratio": "16:9", "size": "1920x1080", "carousel": None,
                  "links_clickable": True, "thumbnail": "1280x720"},
    "facebook":  {"ar": "فيسبوك", "caption": 63206, "title": None, "hashtags": None, "hashtags_tip": "0–3",
                  "visible": 125, "seconds": None, "ratio": "9:16", "size": "1080x1920", "carousel": 10,
                  "links_clickable": True},
    "x":         {"ar": "إكس", "caption": 280, "title": None, "hashtags": None, "hashtags_tip": "0–2",
                  "visible": None, "seconds": 140, "ratio": "16:9", "size": "1280x720", "carousel": 4,
                  "links_clickable": True, "weighted": True},
    "linkedin":  {"ar": "لينكدإن", "caption": 3000, "title": None, "hashtags": None, "hashtags_tip": "3–5",
                  "visible": 210, "seconds": 600, "ratio": "1:1", "size": "1080x1080", "carousel": "pdf",
                  "links_clickable": True},
    "threads":   {"ar": "ثريدز", "caption": 500, "title": None, "hashtags": 1, "hashtags_tip": "1 (وسم موضوع)",
                  "visible": None, "seconds": None, "ratio": "4:5", "size": "1080x1350", "carousel": None,
                  "links_clickable": True},
}

# Starting-point posting slots (local time) when the account has no export yet — a heuristic, never a fact.
DEFAULT_SLOTS = {"tiktok": "21:00", "reels": "20:00", "instagram": "13:00", "shorts": "18:00", "youtube": "17:00",
                 "facebook": "19:00", "x": "12:00", "linkedin": "09:00", "threads": "12:00"}
DAYS = ["mon", "tue", "wed", "thu", "fri", "sat", "sun"]
DAYS_AR = ["الاثنين", "الثلاثاء", "الأربعاء", "الخميس", "الجمعة", "السبت", "الأحد"]

# ---------------------------------------------------------------- text helpers
_TASHKEEL = re.compile(r"[\u064B-\u065F\u0670]")
_URL = re.compile(r"https?://\S+|www\.\S+")
_TAG = re.compile(r"(?<![\w#])#[\w\u0600-\u06FF]+")


def is_arabic(text: str) -> bool:
    return any("\u0600" <= ch <= "\u06FF" for ch in text)


def x_length(text: str) -> int:
    """X's weighted count: URLs are 23, code points in the light ranges are 1, everything else (emoji, CJK) is 2."""
    n = 0
    for part in _URL.split(text):
        for ch in part:
            cp = ord(ch)
            light = cp <= 4351 or 8192 <= cp <= 8205 or 8208 <= cp <= 8223 or 8242 <= cp <= 8247
            n += 1 if light else 2
    return n + 23 * len(_URL.findall(text))


def tag_key(tag: str) -> str:
    t = _TASHKEEL.sub("", tag.lower().lstrip("#")).replace("\u0640", "")
    for a, b in (("أ", "ا"), ("إ", "ا"), ("آ", "ا"), ("ة", "ه"), ("ى", "ي")):
        t = t.replace(a, b)
    return t


def check_tag(tag: str) -> list[str]:
    probs = []
    if not tag.startswith("#"):
        probs.append(f"{tag}: الوسم لازم يبدأ بـ #")
    body = tag.lstrip("#")
    if not body:
        probs.append("وسم فاضي")
    if re.search(r"\s", body):
        probs.append(f"{tag}: فيه مسافة — استعمل _ بدلها")
    if re.search(r"[^\w\u0600-\u06FF]", re.sub(r"[\s_]", "", body)):
        probs.append(f"{tag}: فيه رموز بتقطع الوسم")
    if _TASHKEEL.search(body) or "\u0640" in body:
        probs.append(f"{tag}: فيه تشكيل أو تطويل — الوسم مش هيطابق بحث الناس")
    if body.isdigit():
        probs.append(f"{tag}: أرقام بس — المنصات بتتجاهله")
    return probs


# ---------------------------------------------------------------- probe (ffprobe, else kmotion's tools)
def probe(path: Path) -> dict | None:
    exe = shutil.which("ffprobe")
    if exe:
        r = subprocess.run([exe, "-v", "error", "-select_streams", "v:0", "-show_entries",
                            "stream=width,height:format=duration", "-of", "json", str(path)],
                           capture_output=True, text=True)
        if r.returncode == 0:
            j = json.loads(r.stdout)
            s = (j.get("streams") or [{}])[0]
            return {"w": s.get("width"), "h": s.get("height"), "seconds": float(j["format"]["duration"])}
    km = Path.home() / ".claude" / "skills" / "kosif-montage-motion" / "scripts"
    if km.exists():
        sys.path.insert(0, str(km))
        try:
            import tools  # type: ignore
            i = tools.probe(path)
            return {"w": i["video"]["w"], "h": i["video"]["h"], "seconds": float(i["duration"])}
        except Exception:
            return None
    return None


def _ratio(r: str) -> float:
    a, b = r.split(":")
    return float(a) / float(b)


def fit(platform: str, info: dict | None) -> list[str]:
    if not info or not info.get("w"):
        return []
    p, out = PLATFORMS[platform], []
    if p["seconds"] and info["seconds"] > p["seconds"] + 0.05:
        out.append(f"الفيديو {info['seconds']:.1f} ث أطول من حد {p['ar']} ({p['seconds']} ث): قصّه أو قسّمه")
    have = info["w"] / info["h"]
    if abs(have - _ratio(p["ratio"])) > 0.02:
        out.append(f"النسبة {info['w']}x{info['h']} مش {p['ratio']}: اعمل نسخة بـ kmotion platforms/aspect")
    need = min(int(v) for v in p["size"].split("x"))
    if min(info["w"], info["h"]) < need:
        out.append(f"الدقة {info['w']}x{info['h']} أقل من المطلوب ({p['size']})")
    return out


# ---------------------------------------------------------------- pack + check
def pack(video: Path | None, platforms: list[str], topic: str = "") -> dict:
    info = probe(video) if video else None
    posts = {}
    for name in platforms:
        p = PLATFORMS[name]
        post = {"caption": "", "hashtags": [], "cta": "", "first_comment": "", "alt_text": "", "cover_text": ""}
        if p["title"]:
            post = {"title": "", "description": "", "hashtags": [], "chapters": [], "pinned_comment": ""}
        post["_rules"] = {k: p[k] for k in ("caption", "title", "hashtags", "hashtags_tip", "visible", "seconds", "ratio")}
        post["_file"] = f"{video.stem}_{name}.mp4" if video else ""
        post["_fit"] = fit(name, info)
        posts[name] = post
    return {"kosif_social": 1, "as_of": AS_OF, "topic": topic, "video": str(video) if video else "", "probe": info,
            "language": "ar", "posts": posts}


def check(doc: dict) -> dict:
    info = doc.get("probe")
    if not info and doc.get("video") and Path(doc["video"]).exists():
        info = probe(Path(doc["video"]))
    report, errors, warnings = {}, 0, 0
    for name, post in (doc.get("posts") or {}).items():
        if name not in PLATFORMS:
            report[name] = {"errors": [f"منصة مش معروفة: {name} (المعروف: {', '.join(PLATFORMS)})"], "warnings": []}
            errors += 1
            continue
        p, E, W = PLATFORMS[name], [], []
        tags = [t.strip() for t in post.get("hashtags") or [] if t.strip()]
        body = post.get("description" if p["title"] else "caption", "") or ""
        inline = _TAG.findall(body)
        full = body if not tags else (body.rstrip() + "\n\n" + " ".join(t for t in tags if t not in inline)).strip()
        all_tags = inline + [t for t in tags if t not in inline]
        n = x_length(full) if p.get("weighted") else len(full)
        if p["title"]:
            title = (post.get("title") or "").strip()
            if not title:
                E.append("العنوان فاضي")
            elif len(title) > p["title"]:
                E.append(f"العنوان {len(title)} حرف > {p['title']}")
            elif len(title) > 70:
                W.append(f"العنوان {len(title)} حرف — بيتقطع على الموبايل بعد ~70")
        elif not body.strip():
            E.append("الكابشن فاضي")
        if n > p["caption"]:
            E.append(f"النص {n} حرف > حد {p['ar']} ({p['caption']})")
        if p["hashtags"] is not None and len(all_tags) > p["hashtags"]:
            msg = f"{len(all_tags)} وسم > الحد ({p['hashtags']})"
            if name in ("shorts", "youtube"):
                msg += " — يوتيوب بيتجاهل كل الوسوم"
            E.append(msg)
        elif p["hashtags"] is None and len(all_tags) > 8:
            W.append(f"{len(all_tags)} وسم كتير — المقترح {p['hashtags_tip']}")
        seen = {}
        for t in all_tags:
            E.extend(check_tag(t))
            k = tag_key(t)
            if k in seen:
                W.append(f"وسم مكرر: {seen[k]} و {t}")
            seen[k] = t
        first = next((ln for ln in body.splitlines() if ln.strip()), "")
        if p["visible"] and len(first) > p["visible"]:
            W.append(f"السطر الأول {len(first)} حرف — الخطّاف هيتقطع بعد ~{p['visible']} قبل «المزيد»")
        if first and _TAG.match(first.strip()):
            W.append("الكابشن بيبدأ بوسم — ابدأ بالخطّاف")
        if not p["links_clickable"] and _URL.search(body):
            W.append(f"الروابط مش بتتضغط في {p['ar']} — حطها في البايو أو أول تعليق")
        if not p["title"] and not (post.get("cta") or "").strip() and not re.search(r"(تابع|احفظ|شارك|اكتب|علق|علّق|اطلب|link|follow|save|comment)", body, re.I):
            W.append("مفيش دعوة لفعل (احفظ/تابع/علّق…)")
        if name == "instagram" and not (post.get("alt_text") or "").strip():
            W.append("مفيش نص بديل للصورة (alt text)")
        W.extend(fit(name, info))
        report[name] = {"platform": p["ar"], "length": n, "limit": p["caption"], "hashtags": len(all_tags),
                        "errors": E, "warnings": W, "ok": not E}
        errors += len(E)
        warnings += len(W)
    return {"ok": errors == 0, "errors": errors, "warnings": warnings, "as_of": AS_OF, "platforms": report}


# ---------------------------------------------------------------- fonts and Arabic text
_FONT_DIRS = ["C:/Windows/Fonts", str(Path.home() / "AppData/Local/Microsoft/Windows/Fonts"), "/usr/share/fonts",
              "/usr/local/share/fonts", str(Path.home() / ".fonts"), "/Library/Fonts", "/System/Library/Fonts"]
_BOLD = ["NotoKufiArabic-Bold", "Cairo-Bold", "Tajawal-Bold", "DUBAI-BOLD", "NotoSansArabic-Bold", "segoeuib", "arialbd", "tahomabd", "DejaVuSans-Bold"]
_REG = ["NotoSansArabic-Regular", "Cairo-Regular", "Tajawal-Regular", "DUBAI-REGULAR", "NotoNaskhArabic-Regular", "segoeui", "arial", "tahoma", "DejaVuSans"]


def find_font(names: list[str]) -> str | None:
    files = {}
    for d in _FONT_DIRS:
        dp = Path(d)
        if dp.exists():
            for f in dp.rglob("*"):
                if f.suffix.lower() in (".ttf", ".otf"):
                    files.setdefault(f.stem.lower(), str(f))
    for n in names:
        if n.lower() in files:
            return files[n.lower()]
    return None


def shape(line: str) -> str:
    if not is_arabic(line):
        return line
    try:
        import arabic_reshaper
        from bidi.algorithm import get_display
        return get_display(arabic_reshaper.reshape(line))
    except ImportError:
        return line


def wrap(text: str, font, max_w: float, draw) -> list[str]:
    """Wrap in logical order (then shape each line), so RTL lines keep their reading order."""
    lines = []
    for para in text.split("\n"):
        cur = ""
        for w in para.split():
            cand = (cur + " " + w).strip()
            if not cur or draw.textlength(shape(cand), font=font) <= max_w:
                cur = cand
            else:
                lines.append(cur)
                cur = w
        lines.append(cur)
    return lines


def _hex(c: str) -> tuple[int, int, int]:
    c = c.lstrip("#")
    return tuple(int(c[i:i + 2], 16) for i in (0, 2, 4))  # type: ignore


# ---------------------------------------------------------------- carousel
_MIXED_BOLD = ["segoeuib", "arialbd", "tahomabd", "DejaVuSans-Bold"]
_MIXED_REG = ["segoeui", "arial", "tahoma", "DejaVuSans"]
_CMAPS: dict[str, set | None] = {}


def covers(path: str | None, text: str) -> bool:
    """True when the font has a glyph for every visible character (fontTools; assumed True without it)."""
    if not path:
        return False
    if path not in _CMAPS:
        try:
            from fontTools.ttLib import TTFont
            _CMAPS[path] = set(TTFont(path, lazy=True).getBestCmap())
        except Exception:
            _CMAPS[path] = None
    cm = _CMAPS[path]
    return cm is None or all(ord(ch) in cm for ch in text if not ch.isspace() and unicodedata.category(ch)[0] != "C")


def carousel(spec: dict, out: Path, base: Path) -> dict:
    from PIL import Image, ImageDraw, ImageFont
    W, H = (int(v) for v in str(spec.get("size", "1080x1350")).lower().split("x"))
    brand = spec.get("brand") or {}
    bg, fg = _hex(brand.get("bg", "#0E1116")), _hex(brand.get("fg", "#FFFFFF"))
    accent = _hex(brand.get("accent", "#E7B65A"))
    fb = brand.get("font_bold") or find_font(_BOLD)
    fr = brand.get("font") or find_font(_REG) or fb
    mb, mr = find_font(_MIXED_BOLD) or fb, find_font(_MIXED_REG) or fr
    slides = spec.get("slides") or []
    if not slides:
        raise SystemExit("spec.slides فاضية")
    platform = spec.get("platform", "instagram")
    cap = PLATFORMS.get(platform, {}).get("carousel")
    warnings = []
    if isinstance(cap, int) and len(slides) > cap:
        warnings.append(f"{len(slides)} شريحة > حد {PLATFORMS[platform]['ar']} ({cap})")
    if not fb:
        warnings.append("مفيش خط عربي على الجهاز — النص هيطلع بخط افتراضي")
    out.mkdir(parents=True, exist_ok=True)
    margin, files, images = round(W * 0.08), [], []
    cache: dict = {}

    def font(path, size, text=""):
        if text and not covers(path, shape(text)):    # shaped glyphs (presentation forms), Latin, @handles
            path = mb if path == fb else mr
        key = (path, size)
        if key not in cache:
            cache[key] = ImageFont.truetype(path, size) if path else ImageFont.load_default(size)
        return cache[key]

    def block(text, path, size, min_size, max_lines, max_h, lead, draw):
        """Largest size <= size whose wrapped lines fit max_lines and max_h."""
        while True:
            f = font(path, size, text)
            lines = wrap(text, f, W - 2 * margin, draw)
            if (len(lines) <= max_lines and len(lines) * size * lead <= max_h) or size <= min_size:
                return f, size, lines
            size -= 2

    for i, s in enumerate(slides, 1):
        im = Image.new("RGB", (W, H), bg)
        if s.get("image"):
            ip = Path(s["image"])
            ip = ip if ip.is_absolute() else base / ip
            if ip.exists():
                ph = Image.open(ip).convert("RGB")
                k = max(W / ph.width, H / ph.height)
                ph = ph.resize((round(ph.width * k), round(ph.height * k)), Image.LANCZOS)
                x, y = (ph.width - W) // 2, (ph.height - H) // 2
                im.paste(ph.crop((x, y, x + W, y + H)))
                shade = Image.new("L", (1, H))
                for yy in range(H):
                    shade.putpixel((0, yy), int(255 * (0.35 + 0.5 * (yy / H))))
                im = Image.composite(Image.new("RGB", (W, H), (0, 0, 0)), im, shade.resize((W, H)))
            else:
                warnings.append(f"شريحة {i}: الصورة مش موجودة {ip}")
        draw = ImageDraw.Draw(im)
        title, text = s.get("title", ""), s.get("body", "")
        rtl = is_arabic(title + text)
        xa, anchor = (W - margin, "ra") if rtl else (margin, "la")
        kicker = s.get("kicker") or f"{i}/{len(slides)}"
        ks = round(W * 0.032)
        draw.text((xa, margin), shape(kicker), font=font(fb, ks, kicker), fill=accent, anchor=anchor)
        top, bottom = margin + ks * 2.5, H - margin * 2.4
        avail = bottom - top
        cover = i == 1
        if title:
            tf, ts, tl = block(title, fb, round(W * (0.105 if cover else 0.085)), 40, 4, avail * 0.6, 1.3, draw)
        else:
            tf, ts, tl = None, 0, []
        if len(tl) > 4:
            warnings.append(f"شريحة {i}: العنوان طويل ({len(tl)} سطور) — قصّره")
        th = len(tl) * ts * 1.3
        gap = ts * 0.7 if text and tl else 0
        if text:
            bf, bs, bl = block(text, fr, round(W * (0.05 if cover else 0.054)), 28, 12, avail - th - gap, 1.6, draw)
        else:
            bf, bs, bl = None, 0, []
        bh = len(bl) * bs * 1.6
        if th + gap + bh > avail:
            warnings.append(f"شريحة {i}: النص أطول من الشريحة — قصّره")
        y = top + max(0, (avail - th - gap - bh) * (0.42 if cover else 0.3))
        ink = y
        for ln in tl:
            draw.text((xa, y), shape(ln), font=font(fb, ts, ln), fill=fg, anchor=anchor)
            ink = draw.textbbox((xa, y), shape(ln), font=font(fb, ts, ln), anchor=anchor)[3]
            y += ts * 1.3
        if tl and text:                                  # a short accent rule under the title's real ink
            y = max(y, ink + gap * 0.45)
            x0 = xa - W * 0.12 if rtl else xa
            draw.rectangle((x0, y, x0 + W * 0.12, y + 6), fill=accent)
            y += gap * 0.9
        for ln in bl:
            if y + bs > bottom + bs:
                break
            draw.text((xa, y), shape(ln), font=font(fr, bs, ln), fill=fg, anchor=anchor)
            y += bs * 1.6
        # footer: handle on one side, swipe cue (word + drawn arrow) on the other, progress bar
        fs = round(W * 0.03)
        fy = H - margin
        handle = brand.get("handle", "")
        if handle:
            draw.text((margin if rtl else W - margin, fy), shape(handle), font=font(fr, fs, handle), fill=fg,
                      anchor="ld" if rtl else "rd")
        if i < len(slides):
            cue = spec.get("swipe", "اسحب" if rtl else "Swipe")
            cf = font(fb, fs, cue)
            aw, ah = fs * 1.1, fs * 0.42
            cy = fy - fs * 0.45
            if rtl:   # the next slide sits to the left in RTL reading: word, then an arrow pointing left
                draw.text((xa, fy), shape(cue), font=cf, fill=accent, anchor="rd")
                tx = xa - draw.textlength(shape(cue), font=cf) - fs * 0.4
                draw.line((tx - aw, cy, tx, cy), fill=accent, width=4)
                draw.polygon([(tx - aw - 2, cy), (tx - aw + ah * 1.4, cy - ah), (tx - aw + ah * 1.4, cy + ah)], fill=accent)
            else:
                draw.text((xa, fy), cue, font=cf, fill=accent, anchor="ld")
                tx = xa + draw.textlength(cue, font=cf) + fs * 0.4
                draw.line((tx, cy, tx + aw, cy), fill=accent, width=4)
                draw.polygon([(tx + aw + 2, cy), (tx + aw - ah * 1.4, cy - ah), (tx + aw - ah * 1.4, cy + ah)], fill=accent)
        bar = round(W * i / len(slides))
        draw.rectangle((W - bar, H - 8, W, H) if rtl else (0, H - 8, bar, H), fill=accent)
        f = out / f"slide_{i:02d}.png"
        im.save(f)
        files.append(str(f))
        images.append(im)
    pdf = out / "carousel.pdf"
    images[0].save(pdf, save_all=True, append_images=images[1:], resolution=150)
    cols = min(5, len(images))
    rows = math.ceil(len(images) / cols)
    tw = 300
    th = round(tw * H / W)
    sheet = Image.new("RGB", (cols * tw + (cols + 1) * 10, rows * th + (rows + 1) * 10), (40, 40, 40))
    for k, img in enumerate(images):
        sheet.paste(img.resize((tw, th), Image.LANCZOS), (10 + (k % cols) * (tw + 10), 10 + (k // cols) * (th + 10)))
    sheet.save(out / "contact_sheet.jpg", quality=90)
    return {"slides": files, "pdf": str(pdf), "contact_sheet": str(out / "contact_sheet.jpg"), "size": f"{W}x{H}",
            "font": fb, "warnings": warnings}


# ---------------------------------------------------------------- calendar
def _smooth_wrr(items: list[dict], n: int) -> list[str]:
    """Smooth weighted round-robin: pillars spread evenly by weight, deterministic."""
    cur = [0.0] * len(items)
    total = sum(max(0.0, float(it.get("weight", 1))) for it in items) or 1.0
    out = []
    for _ in range(n):
        for k, it in enumerate(items):
            cur[k] += max(0.0, float(it.get("weight", 1)))
        j = max(range(len(items)), key=lambda k: (cur[k], -k))
        cur[j] -= total
        out.append(items[j]["name"])
    return out


def calendar(plan: dict, out: Path) -> dict:
    start = dt.date.fromisoformat(plan["start"])
    weeks = int(plan.get("weeks", 4))
    off = {d.lower()[:3] for d in plan.get("days_off", [])}
    days = [k for k in range(7) if DAYS[k] not in off] or list(range(7))
    pillars = plan.get("pillars") or [{"name": "عام", "weight": 1}]
    formats = plan.get("formats") or {}
    slots = {**DEFAULT_SLOTS, **(plan.get("slots") or {})}
    ideas = list(plan.get("ideas") or [])
    rows = []
    for name, per_week in (plan.get("platforms") or {}).items():
        if name not in PLATFORMS:
            raise SystemExit(f"منصة مش معروفة: {name}")
        per_week = min(int(per_week), len(days) * 3)
        for w in range(weeks):
            week0 = start + dt.timedelta(days=7 * w)
            for i in range(per_week):
                k = days[(i * len(days)) // per_week % len(days)] if per_week <= len(days) else days[i % len(days)]
                offset = (k - week0.weekday()) % 7
                day = week0 + dt.timedelta(days=offset)
                slot = slots.get(name, "20:00")
                slot = slot[i % len(slot)] if isinstance(slot, list) else slot
                rows.append({"date": day.isoformat(), "time": slot, "platform": name})
    rows.sort(key=lambda r: (r["date"], r["time"], list(PLATFORMS).index(r["platform"])))
    pil = _smooth_wrr(pillars, len(rows))
    fcount: dict[str, int] = {}
    for r, p in zip(rows, pil):
        d = dt.date.fromisoformat(r["date"])
        fl = formats.get(r["platform"]) or []
        j = fcount.get(r["platform"], 0)
        fcount[r["platform"]] = j + 1
        r.update({"day": DAYS_AR[d.weekday()], "platform_ar": PLATFORMS[r["platform"]]["ar"], "pillar": p,
                  "format": fl[j % len(fl)] if fl else "", "idea": ideas.pop(0) if ideas else "", "status": "مسودة"})
    out.mkdir(parents=True, exist_ok=True)
    cols = ["date", "day", "time", "platform", "platform_ar", "pillar", "format", "idea", "status"]
    heads = ["التاريخ", "اليوم", "الوقت", "المنصة", "اسم المنصة", "الركيزة", "الشكل", "الفكرة", "الحالة"]
    with open(out / "calendar.csv", "w", newline="", encoding="utf-8-sig") as fh:
        wr = csv.writer(fh)
        wr.writerow(heads)
        for r in rows:
            wr.writerow([r[c] for c in cols])
    md = ["| " + " | ".join(heads) + " |", "|" + "---|" * len(heads)]
    md += ["| " + " | ".join(str(r[c]) for c in cols) + " |" for r in rows]
    src = "مواعيد من تحليلات الحساب" if plan.get("slots") else "مواعيد مبدئية (افتراض) — بدّلها بنتيجة besttime من تحليلات حسابك"
    (out / "calendar.md").write_text(f"# خطة المحتوى\n\nمن {start} لمدة {weeks} أسابيع — {len(rows)} منشور. {src}.\n\n"
                                     + "\n".join(md) + "\n", encoding="utf-8")
    ics = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//KOSIF Social//AR", "CALSCALE:GREGORIAN"]
    stamp = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    for n, r in enumerate(rows):
        t = r["date"].replace("-", "") + "T" + r["time"].replace(":", "") + "00"
        summary = f"{r['platform_ar']}: {r['idea'] or r['pillar']}".replace(",", "\\,")
        ics += ["BEGIN:VEVENT", f"UID:kosif-social-{t}-{r['platform']}-{n}@local", f"DTSTAMP:{stamp}",
                f"DTSTART:{t}", "DURATION:PT30M", f"SUMMARY:{summary}",
                f"DESCRIPTION:{(r['format'] + ' — ' + r['pillar']).replace(',', '\\,')}", "END:VEVENT"]
    ics.append("END:VCALENDAR")
    (out / "calendar.ics").write_text("\r\n".join(ics) + "\r\n", encoding="utf-8")
    files = [str(out / f) for f in ("calendar.csv", "calendar.md", "calendar.ics")]
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "الخطة"
        ws.sheet_view.rightToLeft = True
        ws.append(heads)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F2937")
        for r in rows:
            ws.append([r[c] for c in cols])
        for col, wdt in zip("ABCDEFGHI", (12, 10, 8, 11, 20, 18, 16, 40, 10)):
            ws.column_dimensions[col].width = wdt
        ws.freeze_panes = "A2"
        wb.save(out / "calendar.xlsx")
        files.append(str(out / "calendar.xlsx"))
    except ImportError:
        pass
    per = {}
    for r in rows:
        per[r["platform"]] = per.get(r["platform"], 0) + 1
    return {"posts": len(rows), "per_platform": per, "slots_source": src, "files": files,
            "pillars": {p["name"]: pil.count(p["name"]) for p in pillars}}


# ---------------------------------------------------------------- best time from the account's own export
_DATE_COLS = ("publish", "posted", "date", "time", "created", "تاريخ", "وقت", "نشر")
_METRIC_COLS = ("views", "view", "plays", "reach", "impressions", "مشاهد", "وصول", "ظهور")


def _parse_dt(v: str) -> dt.datetime | None:
    v = v.strip()
    for f in ("%Y-%m-%dT%H:%M:%S%z", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%m/%d/%Y %H:%M",
              "%d/%m/%Y %H:%M", "%m/%d/%Y %I:%M %p", "%b %d, %Y %I:%M %p", "%Y-%m-%dT%H:%M:%S.%f%z"):
        try:
            return dt.datetime.strptime(v.replace("Z", "+0000"), f)
        except ValueError:
            continue
    try:
        return dt.datetime.fromisoformat(v)
    except ValueError:
        return None


def _num(v: str) -> float | None:
    v = unicodedata.normalize("NFKC", str(v)).replace(",", "").replace("٬", "").strip()
    m = re.fullmatch(r"([\d.]+)\s*([kKmM]?)", v)
    if not m:
        return None
    x = float(m.group(1))
    return x * (1000 if m.group(2).lower() == "k" else 1_000_000 if m.group(2).lower() == "m" else 1)


def besttime(path: Path, metric: str | None = None, date_col: str | None = None, block: int = 3) -> dict:
    with open(path, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    if not rows:
        raise SystemExit("الملف فاضي")
    keys = list(rows[0])
    pick = lambda cands: next((k for c in cands for k in keys if c in k.lower()), None)  # noqa: E731
    dc = date_col or pick(_DATE_COLS)
    mc = metric if metric in keys else (next((k for k in keys if metric and metric.lower() in k.lower()), None) if metric else pick(_METRIC_COLS))
    if not dc or not mc:
        raise SystemExit(f"مش لاقي عمود التاريخ/المقياس. الأعمدة: {keys} — حدد --date-col و --metric")
    cells: dict[tuple[int, int], list[float]] = {}
    used = 0
    for r in rows:
        t, v = _parse_dt(r.get(dc, "")), _num(r.get(mc, ""))
        if t is None or v is None:
            continue
        used += 1
        cells.setdefault((t.weekday(), t.hour // block), []).append(v)
    if used < 8:
        return {"posts_used": used, "verdict": "البيانات قليلة (< 8 منشورات) — استعمل المواعيد المبدئية وجمّع أكتر", "top": []}
    overall = statistics.median(v for vs in cells.values() for v in vs)
    ranked = sorted(((statistics.median(vs), len(vs), k) for k, vs in cells.items() if len(vs) >= 2), reverse=True)
    top = [{"day": DAYS_AR[k[0]], "hours": f"{k[1] * block:02d}:00–{(k[1] + 1) * block:02d}:00", "median": round(m),
            "vs_overall": f"{(m / overall - 1) * 100:+.0f}%" if overall else "", "posts": n} for m, n, k in ranked[:5]]
    return {"posts_used": used, "date_column": dc, "metric": mc, "overall_median": round(overall), "top": top,
            "note": "الوسيط لكل خانة فيها منشورين على الأقل؛ ده ارتباط مش سبب — جرّب الوقت الأعلى أسبوعين وقارن بـ abread."}


# ---------------------------------------------------------------- A/B
def _z(p: float) -> float:
    """Inverse standard normal CDF (Acklam's rational approximation, |error| < 1.2e-9)."""
    a = [-3.969683028665376e+01, 2.209460984245205e+02, -2.759285104469687e+02, 1.383577518672690e+02,
         -3.066479806614716e+01, 2.506628277459239e+00]
    b = [-5.447609879822406e+01, 1.615858368580409e+02, -1.556989798598866e+02, 6.680131188771972e+01,
         -1.328068155288572e+01]
    c = [-7.784894002430293e-03, -3.223964580411365e-01, -2.400758277161838e+00, -2.549732539343734e+00,
         4.374664141464968e+00, 2.938163982698783e+00]
    d = [7.784695709041462e-03, 3.224671290700398e-01, 2.445134137142996e+00, 3.754408661907416e+00]
    lo = 0.02425
    if p < lo:
        q = math.sqrt(-2 * math.log(p))
        return (((((c[0] * q + c[1]) * q + c[2]) * q + c[3]) * q + c[4]) * q + c[5]) / ((((d[0] * q + d[1]) * q + d[2]) * q + d[3]) * q + 1)
    if p > 1 - lo:
        return -_z(1 - p)
    q = p - 0.5
    r = q * q
    return (((((a[0] * r + a[1]) * r + a[2]) * r + a[3]) * r + a[4]) * r + a[5]) * q / (((((b[0] * r + b[1]) * r + b[2]) * r + b[3]) * r + b[4]) * r + 1)


def sample_size(p1: float, lift: float, alpha: float = 0.05, power: float = 0.8) -> int:
    """Impressions per variant to detect p1 → p1·(1+lift) with a two-sided two-proportion test."""
    p2 = p1 * (1 + lift)
    pb = (p1 + p2) / 2
    za, zb = _z(1 - alpha / 2), _z(power)
    n = (za * math.sqrt(2 * pb * (1 - pb)) + zb * math.sqrt(p1 * (1 - p1) + p2 * (1 - p2))) ** 2 / (p2 - p1) ** 2
    return math.ceil(n)


def abplan(spec: dict, out: Path) -> dict:
    p1 = float(spec.get("baseline_rate", 0.03))
    lift = float(spec.get("min_lift", 0.3))
    variants = spec.get("variants") or {}
    if len(variants) < 2:
        raise SystemExit("لازم نسختين على الأقل في variants")
    n = sample_size(p1, lift)
    out.mkdir(parents=True, exist_ok=True)
    with open(out / "ab_results.csv", "w", newline="", encoding="utf-8-sig") as fh:
        wr = csv.writer(fh)
        wr.writerow(["variant", "date", "impressions", "conversions", "note"])
        for k in variants:
            wr.writerow([k, "", "", "", ""])
    plan = {"platform": spec.get("platform", ""), "variable": spec.get("variable", "hook"),
            "metric": spec.get("metric", "نسبة الحفظ أو الضغط"), "variants": variants,
            "baseline_rate": p1, "min_lift": lift, "impressions_per_variant": n,
            "rules": ["غيّر حاجة واحدة بس بين النسخ (الخطّاف أو الغلاف أو العنوان)",
                      "انشر النسخ في نفس الوقت من اليوم وبدّل الترتيب (A ثم B، وبعدها B ثم A)",
                      f"متحكمش قبل ما كل نسخة توصل {n:,} ظهور — النظر المتكرر بيدّي فايز وهمي",
                      "سجّل الأرقام في ab_results.csv وشغّل abread"],
            "tracking_csv": str(out / "ab_results.csv")}
    (out / "ab_plan.json").write_text(json.dumps(plan, ensure_ascii=False, indent=1), encoding="utf-8")
    return plan


def _phi(x: float) -> float:
    return 0.5 * (1 + math.erf(x / math.sqrt(2)))


def abread(path: Path, need: int | None = None) -> dict:
    agg: dict[str, list[float]] = {}
    with open(path, encoding="utf-8-sig", newline="") as fh:
        for r in csv.DictReader(fh):
            i, c = _num(r.get("impressions", "")), _num(r.get("conversions", ""))
            if i is None or c is None:
                continue
            a = agg.setdefault(r["variant"].strip(), [0.0, 0.0])
            a[0] += i
            a[1] += c
    if len(agg) < 2:
        return {"verdict": "محتاج أرقام لنسختين على الأقل", "variants": agg}
    (ka, (na, ca)), (kb, (nb, cb)) = list(agg.items())[:2]
    pa, pb = ca / na, cb / nb
    pool = (ca + cb) / (na + nb)
    se = math.sqrt(pool * (1 - pool) * (1 / na + 1 / nb)) if 0 < pool < 1 else 0
    z = (pb - pa) / se if se else 0.0
    p = 2 * (1 - _phi(abs(z)))
    se_d = math.sqrt(pa * (1 - pa) / na + pb * (1 - pb) / nb)
    lo, hi = (pb - pa) - 1.96 * se_d, (pb - pa) + 1.96 * se_d
    enough = need is None or min(na, nb) >= need
    winner = kb if pb > pa else ka
    if p < 0.05 and enough:
        verdict = f"{winner} أحسن فعلاً (p {'< 0.001' if p < 0.001 else f'= {p:.3f}'})"
    elif p < 0.05:
        verdict = f"{winner} متقدم بس العينة لسه أقل من المطلوب ({need:,}) — كمّل"
    else:
        verdict = f"مفيش فرق مؤكد لسه (p = {p:.2f}) — كمّل أو اعتبرهم متساويين"
    return {"variants": {ka: {"impressions": na, "conversions": ca, "rate": round(pa, 5)},
                         kb: {"impressions": nb, "conversions": cb, "rate": round(pb, 5)}},
            "lift": f"{(pb / pa - 1) * 100:+.1f}%" if pa else "", "diff_95ci": [round(lo, 5), round(hi, 5)],
            "z": round(z, 3), "p_value": round(p, 4), "verdict": verdict}


# ---------------------------------------------------------------- CLI
def _load(p: str) -> dict:
    return json.loads(Path(p).read_text(encoding="utf-8-sig"))


def _print(obj) -> None:
    print(json.dumps(obj, ensure_ascii=False, indent=1))


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("limits"); s.add_argument("--platform")
    s = sub.add_parser("pack"); s.add_argument("video", nargs="?"); s.add_argument("--platforms", default="tiktok,reels,shorts")
    s.add_argument("--topic", default=""); s.add_argument("--out")
    s = sub.add_parser("check"); s.add_argument("pack")
    s = sub.add_parser("carousel"); s.add_argument("spec"); s.add_argument("--out", required=True)
    s = sub.add_parser("calendar"); s.add_argument("plan"); s.add_argument("--out", required=True)
    s = sub.add_parser("besttime"); s.add_argument("csv"); s.add_argument("--metric"); s.add_argument("--date-col")
    s.add_argument("--block", type=int, default=3)
    s = sub.add_parser("abplan"); s.add_argument("spec"); s.add_argument("--out", required=True)
    s = sub.add_parser("abread"); s.add_argument("csv"); s.add_argument("--need", type=int)
    a = ap.parse_args(argv)
    if a.cmd == "limits":
        if a.platform and a.platform not in PLATFORMS:
            print(f"منصة مش معروفة. المعروف: {', '.join(PLATFORMS)}", file=sys.stderr)
            return 2
        _print({"as_of": AS_OF, "platforms": {a.platform: PLATFORMS[a.platform]} if a.platform else PLATFORMS})
        return 0
    if a.cmd == "pack":
        names = [p.strip() for p in a.platforms.split(",") if p.strip()]
        bad = [p for p in names if p not in PLATFORMS]
        if bad:
            print(f"منصات مش معروفة: {bad}. المعروف: {', '.join(PLATFORMS)}", file=sys.stderr)
            return 2
        v = Path(a.video) if a.video else None
        if v and not v.exists():
            print(f"الملف مش موجود: {v}", file=sys.stderr)
            return 2
        doc = pack(v, names, a.topic)
        if a.out:
            Path(a.out).parent.mkdir(parents=True, exist_ok=True)
            Path(a.out).write_text(json.dumps(doc, ensure_ascii=False, indent=1), encoding="utf-8")
            print(a.out)
        else:
            _print(doc)
        return 0
    if a.cmd == "check":
        rep = check(_load(a.pack))
        _print(rep)
        return 0 if rep["ok"] else 1
    if a.cmd == "carousel":
        _print(carousel(_load(a.spec), Path(a.out), Path(a.spec).resolve().parent))
        return 0
    if a.cmd == "calendar":
        _print(calendar(_load(a.plan), Path(a.out)))
        return 0
    if a.cmd == "besttime":
        _print(besttime(Path(a.csv), a.metric, a.date_col, a.block))
        return 0
    if a.cmd == "abplan":
        _print(abplan(_load(a.spec), Path(a.out)))
        return 0
    if a.cmd == "abread":
        _print(abread(Path(a.csv), a.need))
        return 0
    return 2


if __name__ == "__main__":
    sys.exit(main())

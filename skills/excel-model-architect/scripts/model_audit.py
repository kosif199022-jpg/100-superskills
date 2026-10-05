#!/usr/bin/env python3
"""مدقّق النماذج المالية: أرقام مدفونة في الصيغ، ومخالفات اصطلاح الألوان، ودوال متقلبة، ومراجع أوراق مكسورة، وغياب صفوف الفحص، ومفتاح السيناريو، وأخطاء محفوظة.

usage: python model_audit.py model.xlsx [--json]
يحتاج openpyxl. لا يعيد الحساب؛ لفحص القيم شغّل xlsx_recalc.py أولاً.
"""
import argparse
import json
import re
import sys

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("openpyxl غير مثبت: pip install openpyxl")

ERR = {"#REF!", "#N/A", "#VALUE!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!"}
HARDCODE = re.compile(r"(?<![A-Za-z\$!:\.\d])(?:\*|/|\+|-)\s*(\d+\.\d+|\d{2,})(?![\d\.:])")  # e.g. *0.21, *1.05, +5000 ; excludes cell refs like A12, 365 handled below
ALLOWED_NUMBERS = {"365", "12", "100", "1000", "360", "52", "0", "1", "2", "3"}
VOLATILE = re.compile(r"\b(NOW|TODAY|RAND|RANDBETWEEN|OFFSET|INDIRECT)\s*\(", re.I)
SHEET_REF = re.compile(r"(?:'([^']+)'|([A-Za-z_][\w\.]*))!")


def main() -> None:
    ap = argparse.ArgumentParser(); ap.add_argument("file"); ap.add_argument("--json", action="store_true")
    a = ap.parse_args()
    wb = load_workbook(a.file, data_only=False); wv = load_workbook(a.file, data_only=True)
    names = set(wb.sheetnames); findings = []; stats = {"formulas": 0, "inputs": 0, "hardcoded": 0}
    has_checks = any(n.lower() in ("checks", "check", "فحوص") for n in names)
    has_scenario = any(dn.lower() == "scenario" for dn in wb.defined_names)
    for ws in wb.worksheets:
        wsv = wv[ws.title]
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None:
                    continue
                color = (c.font.color.rgb if c.font and c.font.color and c.font.color.type == "rgb" else None) or ""
                is_formula = isinstance(v, str) and v.startswith("=")
                if is_formula:
                    stats["formulas"] += 1
                    f = v.upper()
                    for m in HARDCODE.finditer(v):
                        num = m.group(1)
                        if num not in ALLOWED_NUMBERS:
                            stats["hardcoded"] += 1
                            findings.append({"sev": "HIGH", "cell": f"{ws.title}!{c.coordinate}", "issue": f"رقم مدفون في صيغة: {m.group(0).strip()}", "fix": "انقله إلى خلية مدخل زرقاء في Assumptions وأشر إليها"})
                    if VOLATILE.search(f):
                        findings.append({"sev": "MED", "cell": f"{ws.title}!{c.coordinate}", "issue": "دالة متقلبة (OFFSET/INDIRECT/NOW…)", "fix": "استبدلها بمرجع ثابت أو INDEX"})
                    for m in SHEET_REF.finditer(v):
                        sh = m.group(1) or m.group(2)
                        if sh not in names:
                            findings.append({"sev": "HIGH", "cell": f"{ws.title}!{c.coordinate}", "issue": f"مرجع إلى ورقة غير موجودة: {sh}", "fix": "أصلح الاسم أو أعد الورقة"})
                    if color.endswith("0000FF"):
                        findings.append({"sev": "LOW", "cell": f"{ws.title}!{c.coordinate}", "issue": "صيغة بلون أزرق (الأزرق للمدخلات)", "fix": "اجعلها سوداء"})
                    cached = wsv[c.coordinate].value
                    if isinstance(cached, str) and cached in ERR:
                        findings.append({"sev": "HIGH", "cell": f"{ws.title}!{c.coordinate}", "issue": f"خطأ محفوظ {cached}", "fix": "أصلح الصيغة؛ لا تلفّها بـ IFERROR لإخفائها"})
                elif isinstance(v, (int, float)):
                    stats["inputs"] += 1
                    if ws.title.lower() not in ("assumptions", "inputs", "docs") and not color.endswith("0000FF") and c.row > 3 and c.column > 2:
                        findings.append({"sev": "MED", "cell": f"{ws.title}!{c.coordinate}", "issue": f"رقم ثابت ({v}) خارج ورقة المدخلات وبلون غير أزرق", "fix": "إن كان مدخلاً فانقله إلى Assumptions بالأزرق؛ وإن كان نتيجة فاجعله صيغة"})
    if not has_checks:
        findings.append({"sev": "HIGH", "cell": "-", "issue": "لا ورقة Checks", "fix": "أضف BS_check وCF_check وحالة كلية OK/ERROR"})
    else:
        wk = wv[next(n for n in names if n.lower() in ("checks", "check", "فحوص"))]
        vals = [c.value for row in wk.iter_rows(min_row=4) for c in row if isinstance(c.value, (int, float))]
        if vals and any(abs(x) > 0.01 for x in vals):
            findings.append({"sev": "HIGH", "cell": "Checks", "issue": "قيم فحص غير صفرية محفوظة", "fix": "الميزانية لا تتوازن أو النقد لا يطابق؛ ابحث عن التدفق المفقود، لا Plug"})
        elif not vals:
            findings.append({"sev": "INFO", "cell": "Checks", "issue": "لا قيم محفوظة للفحوص (لم يُعد الحساب)", "fix": "افتح في إكسل أو شغّل xlsx_recalc.py"})
    if not has_scenario:
        findings.append({"sev": "MED", "cell": "-", "issue": "لا اسم معرّف Scenario", "fix": "مفتاح سيناريو واحد يختار أعمدة المحرّكات بـ CHOOSE/INDEX"})
    sev_order = {"HIGH": 0, "MED": 1, "LOW": 2, "INFO": 3}
    findings.sort(key=lambda f: sev_order[f["sev"]])
    verdict = "PASS" if not any(f["sev"] == "HIGH" for f in findings) else "REVISE"
    rep = {"file": a.file, "verdict": verdict, "stats": stats, "findings": findings[:200], "total_findings": len(findings)}
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if a.json:
        print(json.dumps(rep, ensure_ascii=False, indent=1))
    else:
        print(f"{verdict}  {a.file}  formulas={stats['formulas']} inputs={stats['inputs']} hardcoded={stats['hardcoded']}")
        for f in findings[:60]:
            print(f"- [{f['sev']}] {f['cell']}: {f['issue']} → {f['fix']}")
        if len(findings) > 60:
            print(f"… و{len(findings) - 60} أخرى")
    sys.exit(0 if verdict == "PASS" else 1)


if __name__ == "__main__":
    main()
